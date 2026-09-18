from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page setup: narrow margins to fit more on one page ──
section = doc.sections[0]
section.page_width  = Inches(8.5)
section.page_height = Inches(11)
section.top_margin    = Cm(1.2)
section.bottom_margin = Cm(1.2)
section.left_margin   = Cm(1.8)
section.right_margin  = Cm(1.8)

# ── Colour helpers ──
DARK_BLUE  = RGBColor(0x00, 0x47, 0xAB)
MID_BLUE   = RGBColor(0x1e, 0x3a, 0x8a)
LIGHT_BLUE = RGBColor(0x60, 0xa5, 0xfa)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GRAY       = RGBColor(0x60, 0x60, 0x70)
GREEN      = RGBColor(0x16, 0x65, 0x34)
RED        = RGBColor(0x99, 0x10, 0x10)
AMBER      = RGBColor(0x78, 0x35, 0x0f)

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def cell_text(cell, text, size=9, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color

def add_para(doc, text, size=9, bold=False, color=None,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=2, space_after=2):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    return p

def add_heading(doc, text, level_size=11, color=DARK_BLUE, space_before=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(2)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(level_size)
    run.font.bold = True
    run.font.color.rgb = color
    # Bottom border
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'),   'single')
    bottom.set(qn('w:sz'),    '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '0047AB')
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

# ════════════════════════════════════════════════
#  HEADER BANNER
# ════════════════════════════════════════════════
tbl = doc.add_table(rows=1, cols=2)
tbl.style = 'Table Grid'
tbl.autofit = False
tbl.columns[0].width = Inches(4.5)
tbl.columns[1].width = Inches(2.8)

left  = tbl.cell(0,0)
right = tbl.cell(0,1)
set_cell_bg(left,  '0047AB')
set_cell_bg(right, '1e3a8a')

cell_text(left,  'SkyKin Automatic Outbound Dialer', size=16, bold=True, color=WHITE)
left.paragraphs[0].paragraph_format.space_before = Pt(6)
p2 = left.add_paragraph('System Documentation  |  Version 1.0  |  September 2026')
p2.paragraph_format.space_before = Pt(1)
p2.paragraph_format.space_after  = Pt(6)
r2 = p2.runs[0] if p2.runs else p2.add_run('System Documentation  |  Version 1.0  |  September 2026')
r2.font.size  = Pt(8)
r2.font.color.rgb = RGBColor(0xc7,0xd2,0xfe)

cell_text(right, 'Powered by SIP WebRTC\nFusionPBX / FreeSWITCH\nPHP 8 + SQLite Backend',
          size=8, color=RGBColor(0xc7,0xd2,0xfe), align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph()

# ════════════════════════════════════════════════
#  TWO-COLUMN BODY  (manual table trick)
# ════════════════════════════════════════════════
body = doc.add_table(rows=1, cols=2)
body.style = 'Table Grid'
body.autofit = False
body.columns[0].width = Inches(3.55)
body.columns[1].width = Inches(3.55)
L = body.cell(0,0)
R = body.cell(0,1)
# Remove borders on body table
for cell in [L, R]:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'none')
        tcBorders.append(el)
    tcPr.append(tcBorders)

# ─── LEFT COLUMN ───────────────────────────────

def lp(text, size=9, bold=False, color=None, space_before=1, space_after=1):
    p = L.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = color
    return p

def lh(text):
    p = lp(text, size=10, bold=True, color=DARK_BLUE, space_before=6, space_after=2)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '4')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '0047AB')
    pBdr.append(bot)
    pPr.append(pBdr)

lh('1. What Is the Automatic Dialer?')
lp('A web-based call center tool that automatically dials customers '
   'at scheduled dates and times. It connects to FusionPBX via SIP WebRTC '
   'and plays a pre-recorded IVR audio when the customer answers.')

lh('2. How to Start the System')
lp('Step 1 — Start the PHP Server', bold=True)
lp('Double-click start_dialer.bat  OR  run in terminal:')
lp('  php.exe -S 0.0.0.0:8080 -t .', size=8, color=RGBColor(0x10,0x40,0x80))
lp('Step 2 — Open the Dashboard', bold=True)
lp('Open browser and go to:  http://localhost:8080')
lp('Step 3 — Log In', bold=True)
lp('Enter your username and password on the login page.')

lh('3. Adding Contacts')
lp('Option A — Upload Excel / CSV File', bold=True)
lp('Click Upload File tab, select your .xlsx or .csv file, click Upload & Import.')
lp('Option B — Manual Entry', bold=True)
lp('Type one per line:  Name, Phone, Date & Time, Notes')
lp('  e.g. Hawi Tadesse, +251939777880, Sep 16 2026 2:00 PM, Account', size=8, color=GRAY)
lp('Option C — Quick Add Form', bold=True)
lp('Fill in the Quick Add form at the top of the queue section.')

lh('4. Excel File Format  (4 Columns)')
# Mini table
t = L.add_table(rows=5, cols=4)
t.style = 'Table Grid'
hdr_row = t.rows[0]
hdrs = ['Col A\nName','Col B\nPhone','Col C\nDate & Time','Col D\nNotes']
hdr_colors = ['0047AB','0047AB','0047AB','0047AB']
for ci, (h, hc) in enumerate(zip(hdrs, hdr_colors)):
    set_cell_bg(hdr_row.cells[ci], hc)
    cell_text(hdr_row.cells[ci], h, size=7, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)

sample = [
    ['Hawi Tadesse', '+251939777880', 'Sep 16 2026 02:00PM', 'Account Verify'],
    ['Abebe Bekele', '+251911234567', 'Sep 16 2026 03:30PM', 'Service Update'],
    ['Sara Ahmed',   '+251922345678', 'Sep 17 2026 10:00AM', 'Follow-up'],
    ['Dawit Haile',  '+251933456789', 'Immediate',            'Priority Now'],
]
for ri, row_data in enumerate(sample):
    row = t.rows[ri+1]
    rbg = 'dbeafe' if ri % 2 == 0 else 'eff6ff'
    for ci, val in enumerate(row_data):
        set_cell_bg(row.cells[ci], rbg)
        fc = RGBColor(0x78,0x35,0x0f) if val == 'Immediate' else RGBColor(0x1e,0x29,0x3b)
        cell_text(row.cells[ci], val, size=7, color=fc)

lp('')
lp('Tip: Write "Immediate" to call the contact as soon as the dialer starts.', size=7.5, color=GRAY)

lh('5. SIP Phone Settings')
t2 = L.add_table(rows=6, cols=2)
t2.style = 'Table Grid'
set_cell_bg(t2.rows[0].cells[0], '0047AB')
set_cell_bg(t2.rows[0].cells[1], '0047AB')
cell_text(t2.rows[0].cells[0], 'Field',   size=7.5, bold=True, color=WHITE)
cell_text(t2.rows[0].cells[1], 'Example', size=7.5, bold=True, color=WHITE)
sip_rows = [
    ('SIP Server / IP', '192.168.1.100'),
    ('Port (WebSocket)', '7443'),
    ('Extension', '101'),
    ('Password', 'your_password'),
    ('Domain / Realm', '192.168.1.100'),
]
for ri, (f, ex) in enumerate(sip_rows):
    rbg = 'dbeafe' if ri%2==0 else 'eff6ff'
    set_cell_bg(t2.rows[ri+1].cells[0], rbg)
    set_cell_bg(t2.rows[ri+1].cells[1], rbg)
    cell_text(t2.rows[ri+1].cells[0], f,  size=7.5, bold=True)
    cell_text(t2.rows[ri+1].cells[1], ex, size=7.5)

# ─── RIGHT COLUMN ──────────────────────────────

def rp(text, size=9, bold=False, color=None, space_before=1, space_after=1):
    p = R.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = color
    return p

def rh(text):
    p = rp(text, size=10, bold=True, color=DARK_BLUE, space_before=6, space_after=2)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot  = OxmlElement('w:bottom')
    bot.set(qn('w:val'),   'single')
    bot.set(qn('w:sz'),    '4')
    bot.set(qn('w:space'), '1')
    bot.set(qn('w:color'), '0047AB')
    pBdr.append(bot)
    pPr.append(pBdr)

rh('6. Running the Auto-Dialer')
steps = [
    ('1', 'Check SIP is Registered',       'Make sure the status shows green (Registered).'),
    ('2', 'Verify Queue',                   'Contacts must be in Pending status with a future scheduled time.'),
    ('3', 'Select IVR Audio',               'In Settings, pick the audio file to play when customer answers.'),
    ('4', 'Click "Start Auto-Dialer"',      'The dialer finds the next scheduled contact and shows a live countdown.'),
    ('5', 'Automatic Call at Scheduled Time','When the time arrives the call is placed automatically via SIP.'),
    ('6', 'Result is Logged',               'Completed / No Answer / Failed — all saved in the call log.'),
]
for num, title, desc in steps:
    p = R.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(1)
    r1 = p.add_run(f'  {num}. ')
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = DARK_BLUE
    r2 = p.add_run(title + ' — ')
    r2.font.bold = True
    r2.font.size = Pt(8.5)
    r3 = p.add_run(desc)
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = GRAY

rh('7. Call Status Guide')
t3 = R.add_table(rows=6, cols=3)
t3.style = 'Table Grid'
for ci,h in enumerate(['Status','Meaning','Action']):
    set_cell_bg(t3.rows[0].cells[ci], '0047AB')
    cell_text(t3.rows[0].cells[ci], h, size=7.5, bold=True, color=WHITE)

status_data = [
    ('Pending',   'dbeafe', 'Waiting for scheduled time',         'Dialer calls automatically'),
    ('Calling',   'fef9c3', 'Call is being placed / ringing',     'Wait for answer'),
    ('Completed', 'd1fae5', 'Customer answered, IVR played',      'Done — logged'),
    ('No Answer', 'fee2e2', 'Customer did not pick up',           'Click Requeue to retry'),
    ('Failed',    'f1f5f9', 'SIP error / not registered',         'Check SIP Settings'),
]
for ri, (status, rbg, meaning, action) in enumerate(status_data):
    set_cell_bg(t3.rows[ri+1].cells[0], rbg)
    set_cell_bg(t3.rows[ri+1].cells[1], rbg)
    set_cell_bg(t3.rows[ri+1].cells[2], rbg)
    cell_text(t3.rows[ri+1].cells[0], status,  size=7.5, bold=True)
    cell_text(t3.rows[ri+1].cells[1], meaning, size=7.5)
    cell_text(t3.rows[ri+1].cells[2], action,  size=7.5)

rh('8. Troubleshooting')
issues = [
    ('SIP shows Not Registered',
     'Check server IP, port, extension, and password in Settings. Click Connect.'),
    ('Call shows Failed',
     'Re-click Connect in Settings. Verify phone number has country code (+251...).'),
    ('Dialer does not call on time',
     'Ensure Auto-Dialer is started, contact is Pending, and scheduled time is in the future.'),
    ('No contacts after Excel upload',
     'Row 1 must be header. Check phone numbers are valid. Try uploading as CSV instead.'),
]
for title, desc in issues:
    p = R.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after  = Pt(2)
    r1 = p.add_run(title + ': ')
    r1.font.bold = True
    r1.font.size = Pt(8.5)
    r1.font.color.rgb = RGBColor(0x7f,0x1d,0x1d)
    r2 = p.add_run(desc)
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GRAY

rh('9. System Architecture')
rp('Browser (Dashboard + SIP.js WebRTC softphone)', size=8, bold=True)
rp('       |  HTTP API calls          |  WSS SIP signaling', size=7.5, color=GRAY)
rp('PHP Backend (api.php + SQLite)    FusionPBX / FreeSWITCH', size=8, bold=True)
rp('                                         |  PSTN Trunk', size=7.5, color=GRAY)
rp('                               Customer Mobile / Landline', size=8, bold=True)
rp('')
rp('Key Files:', size=8, bold=True, color=DARK_BLUE)
files = [
    ('index.php',           'Main dashboard page'),
    ('api.php',             'Backend API and database logic'),
    ('data/dialer.db',      'SQLite database — all leads and logs'),
    ('assets/js/dialer.js', 'Core auto-dialer JavaScript engine'),
    ('assets/audio/',       'Uploaded IVR audio recordings'),
    ('start_dialer.bat',    'One-click server launch script'),
]
for fn, desc in files:
    p = R.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(1)
    r1 = p.add_run(f'  {fn}')
    r1.font.size = Pt(7.5)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x10,0x40,0x80)
    r2 = p.add_run(f'  —  {desc}')
    r2.font.size = Pt(7.5)
    r2.font.color.rgb = GRAY

# ── Footer ──
doc.add_paragraph()
fp = doc.add_paragraph('SkyKin Solutions  |  Automatic Outbound Dialer  |  2026  |  http://localhost:8080')
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(4)
fr = fp.runs[0]
fr.font.size = Pt(7.5)
fr.font.color.rgb = GRAY
fr.font.italic = True

out = r'C:\Users\user\Desktop\call center test\automatic dialer\SkyKin_AutoDialer_Documentation.docx'
doc.save(out)
print('Saved:', out)
