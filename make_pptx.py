from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

# Professional palette — navy + white + light grey only
NAVY      = RGBColor(0x1B, 0x2A, 0x4A)   # dark navy
NAVY2     = RGBColor(0x23, 0x3D, 0x6E)   # mid navy
STEEL     = RGBColor(0x4A, 0x6F, 0xA5)   # steel blue accent
SILVER    = RGBColor(0xF2, 0xF4, 0xF8)   # very light grey bg
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
CHARCOAL  = RGBColor(0x2D, 0x2D, 0x2D)
MIDGRAY   = RGBColor(0x6B, 0x7B, 0x8D)
LIGHTGRAY = RGBColor(0xD0, 0xD7, 0xE2)

BLANK = prs.slide_layouts[6]

def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color

def rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s

def txt(slide, text, l, t, w, h, size=12, bold=False,
        color=CHARCOAL, align=PP_ALIGN.LEFT, wrap=True, italic=False):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    p  = tf.paragraphs[0]
    p.alignment = align
    r  = p.add_run()
    r.text = text
    r.font.size   = Pt(size)
    r.font.bold   = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name   = 'Calibri'
    return tb

def divider(slide, y):
    rect(slide, 0.55, y, 12.23, 0.025, LIGHTGRAY)

# ════════════════════════════════════
# SLIDE 1  — Title
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 0.18, 7.5, NAVY)          # left bar
rect(sl, 0.18, 0, 5.2, 7.5, NAVY)        # left panel
rect(sl, 0.18, 5.8, 5.2, 0.06, STEEL)    # accent line

txt(sl, 'SKYKIN SOLUTIONS', 0.55, 1.6, 4.5, 0.5,
    size=10, bold=True, color=LIGHTGRAY, italic=True)
txt(sl, 'Automatic Outbound', 0.55, 2.15, 4.7, 1.0,
    size=36, bold=True, color=WHITE)
txt(sl, 'Dialer', 0.55, 3.05, 4.7, 0.9,
    size=36, bold=True, color=RGBColor(0xA8,0xBE,0xDC))
txt(sl, 'Intelligent scheduled call management system\npowered by SIP / WebRTC technology',
    0.55, 4.05, 4.7, 0.9, size=12, color=LIGHTGRAY, wrap=True)
txt(sl, 'September 2026  |  Version 1.0', 0.55, 6.1, 4.5, 0.45,
    size=10, color=MIDGRAY, italic=True)

# Right side — 3 key facts
facts = [
    ('Web-Based',       'Runs entirely in the browser.\nNo software installation required.'),
    ('SIP / WebRTC',    'Real phone calls via FusionPBX\nover a secure WebSocket connection.'),
    ('Fully Automated', 'Contacts are called automatically\nat their exact scheduled date and time.'),
]
for i,(title,desc) in enumerate(facts):
    y = 1.5 + i*1.85
    rect(sl, 6.0, y, 6.8, 1.6, SILVER)
    rect(sl, 6.0, y, 0.07, 1.6, STEEL)
    txt(sl, title, 6.2, y+0.15, 6.3, 0.45, size=14, bold=True, color=NAVY)
    txt(sl, desc,  6.2, y+0.6,  6.3, 0.9,  size=11, color=MIDGRAY, wrap=True)

# ════════════════════════════════════
# SLIDE 2  — Agenda
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 13.33, 1.1, NAVY)
rect(sl, 0, 1.1, 13.33, 0.06, STEEL)
txt(sl, 'Agenda', 0.55, 0.22, 12, 0.65, size=28, bold=True, color=WHITE)

items = [
    ('01', 'Overview',          'What the Automatic Dialer is and what problem it solves'),
    ('02', 'How It Works',      'Step-by-step workflow from upload to completed call'),
    ('03', 'Key Features',      'Core capabilities of the system'),
    ('04', 'Call Statuses',     'Understanding each possible call outcome'),
    ('05', 'File Import Format','How to prepare your Excel contact list'),
    ('06', 'System Architecture','Technical components and data flow'),
]
for i,(num,title,desc) in enumerate(items):
    col = i % 2
    row = i // 2
    x = 0.55 + col*6.5
    y = 1.5  + row*1.7
    rect(sl, x, y, 6.0, 1.5, SILVER)
    txt(sl, num,   x+0.2,  y+0.15, 0.8, 0.55, size=20, bold=True, color=STEEL)
    txt(sl, title, x+0.2,  y+0.62, 5.5, 0.45, size=13, bold=True, color=NAVY)
    txt(sl, desc,  x+0.2,  y+1.05, 5.5, 0.4,  size=10, color=MIDGRAY, wrap=True)

# ════════════════════════════════════
# SLIDE 3  — Overview
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 13.33, 1.1, NAVY)
rect(sl, 0, 1.1, 13.33, 0.06, STEEL)
txt(sl, '01  —  Overview', 0.55, 0.22, 12, 0.65, size=28, bold=True, color=WHITE)

txt(sl,
    'The SkyKin Automatic Outbound Dialer is a web-based system that reads a list of '
    'customer contacts and automatically places phone calls at their scheduled date and time. '
    'When a customer answers, a pre-recorded IVR message is played. Every call result is '
    'logged for review.',
    0.55, 1.35, 12.2, 1.0, size=12, color=CHARCOAL, wrap=True)

divider(sl, 2.55)

points = [
    ('No Manual Dialing',   'Agents do not need to dial numbers manually. The system handles all outbound calls automatically.'),
    ('Time-Based Execution','Each contact is assigned a specific date and time. The dialer waits and calls precisely on schedule.'),
    ('IVR Integration',     'A pre-recorded audio message is automatically played to the customer upon answering the call.'),
    ('Full Audit Trail',    'Every call attempt is recorded with status, timestamp, duration, and the audio file played.'),
]
for i,(title,desc) in enumerate(points):
    col = i % 2
    row = i // 2
    x = 0.55 + col*6.45
    y = 2.75  + row*2.1
    rect(sl, x, y, 6.1, 1.9, SILVER)
    rect(sl, x, y, 6.1, 0.06, STEEL)
    txt(sl, title, x+0.25, y+0.2,  5.6, 0.45, size=13, bold=True, color=NAVY)
    txt(sl, desc,  x+0.25, y+0.7,  5.6, 1.1,  size=11, color=CHARCOAL, wrap=True)

# ════════════════════════════════════
# SLIDE 4  — How It Works
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 13.33, 1.1, NAVY)
rect(sl, 0, 1.1, 13.33, 0.06, STEEL)
txt(sl, '02  —  How It Works', 0.55, 0.22, 12, 0.65, size=28, bold=True, color=WHITE)

steps = [
    ('Upload Contact List',
     'Import an Excel (.xlsx) or CSV file containing customer names, phone numbers, '
     'scheduled call date and time, and optional notes.'),
    ('Start the Auto-Dialer',
     'Select the IVR audio recording, configure the pacing delay, then click '
     'Start Auto-Dialer. The system begins monitoring the schedule queue every second.'),
    ('Automatic Call at Scheduled Time',
     'When the current time matches a contact\'s scheduled time, the dialer automatically '
     'places a SIP call. A live countdown timer shows the time remaining until the next call.'),
    ('Result Logged',
     'The call outcome (Completed, No Answer, or Failed) is recorded along with the timestamp, '
     'duration, and IVR audio used. Results are visible in the dashboard call log.'),
]
for i,(title,desc) in enumerate(steps):
    y = 1.45 + i*1.42
    rect(sl, 0.55, y, 12.2, 1.28, SILVER)
    rect(sl, 0.55, y, 0.07, 1.28, NAVY)
    txt(sl, str(i+1), 0.75, y+0.25, 0.55, 0.65, size=22, bold=True, color=STEEL)
    txt(sl, title,    1.4,  y+0.1,  10.8, 0.45, size=13, bold=True, color=NAVY)
    txt(sl, desc,     1.4,  y+0.58, 10.8, 0.65, size=11, color=CHARCOAL, wrap=True)

# ════════════════════════════════════
# SLIDE 5  — Key Features
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 13.33, 1.1, NAVY)
rect(sl, 0, 1.1, 13.33, 0.06, STEEL)
txt(sl, '03  —  Key Features', 0.55, 0.22, 12, 0.65, size=28, bold=True, color=WHITE)

features = [
    ('Date and Time Scheduling',
     'Schedule calls for any specific future date and time. The dialer calls precisely at the moment scheduled.'),
    ('Live Countdown Display',
     'The dashboard shows a real-time countdown to the next scheduled call so operators know exactly when it will fire.'),
    ('Excel and CSV Import',
     'Bulk-import hundreds of contacts at once using a standard Excel or CSV file. Headers are auto-detected.'),
    ('Real SIP Phone Calls',
     'Calls are placed via WebRTC through FusionPBX to real customer mobile or landline phones over PSTN.'),
    ('IVR Audio Playback',
     'A pre-recorded greeting message is automatically streamed into the answered call without agent involvement.'),
    ('Full Call Logs',
     'Every call is logged with status, timestamp, duration, and audio file. Logs can be exported to CSV.'),
]
for i,(title,desc) in enumerate(features):
    col = i % 2
    row = i // 2
    x = 0.55 + col*6.45
    y = 1.35  + row*1.92
    rect(sl, x, y, 6.1, 1.78, SILVER)
    rect(sl, x, y, 0.07, 1.78, STEEL)
    txt(sl, title, x+0.25, y+0.15, 5.6, 0.45, size=13, bold=True, color=NAVY)
    txt(sl, desc,  x+0.25, y+0.62, 5.6, 1.05, size=11, color=CHARCOAL, wrap=True)

# ════════════════════════════════════
# SLIDE 6  — Call Statuses
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 13.33, 1.1, NAVY)
rect(sl, 0, 1.1, 13.33, 0.06, STEEL)
txt(sl, '04  —  Call Statuses', 0.55, 0.22, 12, 0.65, size=28, bold=True, color=WHITE)

txt(sl, 'Each contact in the queue is assigned one of the following statuses throughout the dialing process:',
    0.55, 1.3, 12.2, 0.45, size=12, color=MIDGRAY, wrap=True)

statuses = [
    ('Pending',   'The contact is waiting for its scheduled date and time to arrive. No action has been taken yet.',     'Dialer will call automatically at the scheduled time.'),
    ('Calling',   'The call has been placed and is currently dialing or ringing the customer.',                          'System is waiting for the customer to answer.'),
    ('Completed', 'The customer answered the call and the IVR audio message was played successfully.',                   'Call is fully logged. The dialer moves to the next contact.'),
    ('No Answer', 'The customer did not answer, the line was busy, or the call was rejected.',                           'Click Requeue to set it back to Pending and retry later.'),
    ('Failed',    'A SIP or network error prevented the call from being placed.',                                        'Check SIP Settings and verify the PBX connection is active.'),
]
for i,(status,meaning,action) in enumerate(statuses):
    y = 1.92 + i*1.05
    rect(sl, 0.55, y, 12.2, 0.95, SILVER if i%2==0 else WHITE)
    rect(sl, 0.55, y, 0.07, 0.95, NAVY)
    txt(sl, status,  0.75, y+0.08, 1.7, 0.4, size=13, bold=True, color=NAVY)
    txt(sl, meaning, 2.6,  y+0.08, 6.0, 0.8, size=11, color=CHARCOAL, wrap=True)
    txt(sl, 'Action: '+action, 8.75, y+0.08, 3.9, 0.8, size=10, color=MIDGRAY, italic=True, wrap=True)

# ════════════════════════════════════
# SLIDE 7  — Excel Format
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 13.33, 1.1, NAVY)
rect(sl, 0, 1.1, 13.33, 0.06, STEEL)
txt(sl, '05  —  Excel Import File Format', 0.55, 0.22, 12, 0.65, size=28, bold=True, color=WHITE)

txt(sl, 'Your Excel or CSV file must follow this 4-column structure. The first row is treated as a header and is skipped automatically.',
    0.55, 1.3, 12.2, 0.5, size=12, color=MIDGRAY, wrap=True)

# Header row
col_x = [0.55, 3.5, 6.45, 10.15]
col_w = [2.85, 2.85, 3.6, 3.0]
col_h = ['Column A  —  Customer Name', 'Column B  —  Phone Number',
         'Column C  —  Scheduled Date & Time', 'Column D  —  Notes']
for x,w,h in zip(col_x, col_w, col_h):
    rect(sl, x, 1.95, w-0.05, 0.58, NAVY)
    txt(sl, h, x+0.12, 2.0, w-0.2, 0.48, size=10, bold=True, color=WHITE, wrap=True)

rows = [
    ['Hawi Tadesse',  '+251939777880', 'Sep 16, 2026  02:00 PM', 'Account Verification'],
    ['Abebe Bekele',  '+251911234567', 'Sep 16, 2026  03:30 PM', 'Service Update'],
    ['Sara Ahmed',    '+251922345678', 'Sep 17, 2026  10:00 AM', 'Follow-up Call'],
    ['Dawit Haile',   '+251933456789', 'Immediate',               'Priority — Call Now'],
    ['Tigist Alemu',  '+251944567890', 'Sep 18, 2026  09:00 AM', 'New Customer Welcome'],
]
for ri,row in enumerate(rows):
    rbg = SILVER if ri%2==0 else WHITE
    for x,w,val in zip(col_x, col_w, row):
        rect(sl, x, 2.58+ri*0.78, w-0.05, 0.72, rbg)
        rect(sl, x, 2.58+ri*0.78, 0.05, 0.72, LIGHTGRAY)
        fc = STEEL if val == 'Immediate' else CHARCOAL
        txt(sl, val, x+0.12, 2.62+ri*0.78, w-0.2, 0.64, size=10.5, color=fc, wrap=True)

txt(sl, 'Note: Write "Immediate" in Column C to call the contact as soon as the Auto-Dialer starts.',
    0.55, 7.0, 12.2, 0.38, size=10, color=MIDGRAY, italic=True)

# ════════════════════════════════════
# SLIDE 8  — Architecture
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 13.33, 1.1, NAVY)
rect(sl, 0, 1.1, 13.33, 0.06, STEEL)
txt(sl, '06  —  System Architecture', 0.55, 0.22, 12, 0.65, size=28, bold=True, color=WHITE)

layers = [
    ('Browser\n(Frontend)',
     'HTML + CSS + JavaScript dashboard.\nSIP.js WebRTC softphone embedded.\nNo agent software installation required.'),
    ('PHP API\n(Backend)',
     'api.php handles all data operations.\nSQLite database stores leads and logs.\nREST API for frontend communication.'),
    ('FusionPBX\n(PBX Server)',
     'SIP registrar for extensions.\nRoutes calls to PSTN trunk.\nSecure WebSocket (WSS) for browser SIP.'),
    ('Customer\n(PSTN)',
     'Receives the outbound call on\ntheir mobile or landline phone.\nHears the IVR audio message.'),
]
for i,(title,desc) in enumerate(layers):
    x = 0.55 + i*3.18
    rect(sl, x, 1.45, 3.0, 4.4, SILVER)
    rect(sl, x, 1.45, 3.0, 0.06, NAVY)
    txt(sl, title, x+0.15, 1.6,  2.7, 0.75, size=14, bold=True, color=NAVY, wrap=True, align=PP_ALIGN.CENTER)
    txt(sl, desc,  x+0.15, 2.5,  2.7, 2.8,  size=11, color=CHARCOAL, wrap=True, align=PP_ALIGN.CENTER)
    if i < 3:
        txt(sl, '→', x+3.0, 3.3, 0.2, 0.5, size=18, bold=True, color=STEEL, align=PP_ALIGN.CENTER)

rect(sl, 0.55, 6.1, 12.2, 0.06, LIGHTGRAY)
txt(sl, 'Data Flow:  Browser  →  WSS (Secure WebSocket)  →  FusionPBX  →  PSTN Trunk  →  Customer Phone',
    0.55, 6.25, 12.2, 0.45, size=11, color=MIDGRAY, italic=True, align=PP_ALIGN.CENTER)

# ════════════════════════════════════
# SLIDE 9  — Closing
# ════════════════════════════════════
sl = prs.slides.add_slide(BLANK)
bg(sl, WHITE)
rect(sl, 0, 0, 0.18, 7.5, NAVY)
rect(sl, 0.18, 0, 5.2, 7.5, NAVY)
rect(sl, 0.18, 5.8, 5.2, 0.06, STEEL)

txt(sl, 'SKYKIN SOLUTIONS', 0.55, 1.8, 4.5, 0.5, size=10, bold=True, color=LIGHTGRAY, italic=True)
txt(sl, 'Thank You', 0.55, 2.35, 4.7, 0.85, size=38, bold=True, color=WHITE)
txt(sl, 'Ready to automate\nyour outbound calls.', 0.55, 3.3, 4.7, 1.0, size=16, color=RGBColor(0xA8,0xBE,0xDC), wrap=True)
txt(sl, 'Dashboard:  http://localhost:8080', 0.55, 4.6, 4.7, 0.5, size=11, color=LIGHTGRAY, italic=True)

summary = [
    ('Automated',  'No manual dialing — calls placed automatically'),
    ('Scheduled',  'Exact date and time per contact'),
    ('Integrated', 'Real calls via FusionPBX over SIP/WebRTC'),
    ('Logged',     'Full call history with status and duration'),
]
for i,(title,desc) in enumerate(summary):
    y = 1.5 + i*1.4
    rect(sl, 6.0, y, 6.8, 1.25, SILVER)
    rect(sl, 6.0, y, 0.07, 1.25, STEEL)
    txt(sl, title, 6.2, y+0.12, 6.3, 0.42, size=14, bold=True, color=NAVY)
    txt(sl, desc,  6.2, y+0.58, 6.3, 0.6,  size=11, color=MIDGRAY, wrap=True)

out = r'C:\Users\user\Desktop\SkyKin_AutoDialer_Presentation.pptx'
prs.save(out)
print('Saved:', out)
