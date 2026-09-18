import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from docx import Document
from docx.shared import Pt as DocxPt, Inches as DocxInches, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ==========================================
# 1. GENERATE PROFESSIONAL PPTX
# ==========================================
def build_pptx(output_paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Corporate Monochrome/Navy Palette
    BG_LIGHT = PptxRGBColor(248, 249, 250)
    WHITE = PptxRGBColor(255, 255, 255)
    NAVY_DARK = PptxRGBColor(15, 23, 42)      # #0F172A
    NAVY_PRIMARY = PptxRGBColor(30, 58, 138)  # #1E3A8A
    TEXT_MAIN = PptxRGBColor(30, 41, 59)      # #1E293B
    TEXT_MUTED = PptxRGBColor(100, 116, 139)  # #64748B
    BORDER_COLOR = PptxRGBColor(226, 232, 240)# #E2E8F0
    CARD_BG = PptxRGBColor(255, 255, 255)

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, category, title, dark_mode=False):
        # Category / Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.5), Inches(11.5), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = PptxRGBColor(148, 163, 184) if dark_mode else NAVY_PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.85), Inches(11.5), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE if dark_mode else NAVY_DARK

        # Subtle divider
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.5), Inches(11.533), Inches(0.015))
        div.fill.solid()
        div.fill.fore_color.rgb = PptxRGBColor(51, 65, 85) if dark_mode else BORDER_COLOR
        div.line.fill.background()

    def add_card(slide, left, top, width, height, title, points, bg_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.25), width - Inches(0.5), height - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY_DARK
        p.space_after = Pt(10)

        for pt_text in points:
            p2 = tf.add_paragraph()
            p2.text = f"•  {pt_text}"
            p2.font.size = Pt(11)
            p2.font.color.rgb = TEXT_MAIN
            p2.space_after = Pt(6)

    # ------------------------------------------
    # SLIDE 1: Title Slide (Executive Navy Dark)
    # ------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, NAVY_DARK)

    # Accent line
    top_line = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(1.0), Inches(0.06))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = PptxRGBColor(59, 130, 246)
    top_line.line.fill.background()

    # Title & Subtitle
    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(10.5), Inches(3.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "SkyKin Automatic Outbound Dialer"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "Automated SIP/WebRTC Telephony Dispatching & IVR Broadcast Platform"
    p2.font.size = Pt(18)
    p2.font.color.rgb = PptxRGBColor(203, 213, 225)
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "Enterprise Solution Overview & Technical Specification"
    p3.font.size = Pt(13)
    p3.font.color.rgb = PptxRGBColor(148, 163, 184)

    meta_box = s1.shapes.add_textbox(Inches(1.2), Inches(5.8), Inches(10.5), Inches(0.8))
    tf_m = meta_box.text_frame
    pm = tf_m.paragraphs[0]
    pm.text = "SkyKin Solutions  |  System Documentation & Executive Presentation  |  Version 1.0"
    pm.font.size = Pt(11)
    pm.font.color.rgb = PptxRGBColor(100, 116, 139)

    # ------------------------------------------
    # SLIDE 2: Executive Summary & Objective
    # ------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, BG_LIGHT)
    add_header(s2, "Overview", "Executive Summary & System Purpose")

    card_w = Inches(3.6)
    card_h = Inches(4.9)
    top_y = Inches(1.8)

    add_card(s2, Inches(0.9), top_y, card_w, card_h, "1. Core Objective", [
        "Automate high-volume outbound calls without manual agent intervention.",
        "Deliver clear pre-recorded voice broadcasts (IVR) upon call pickup.",
        "Ensure consistent operational reach and follow-up reliability."
    ])

    add_card(s2, Inches(4.85), top_y, card_w, card_h, "2. Key Capabilities", [
        "Dynamic date & time scheduled dispatching.",
        "Real-time countdown and closest schedule detection.",
        "Batch Excel (.xlsx / .csv) contact list ingestion.",
        "Live SIP telephony integration via secure WebSockets."
    ])

    add_card(s2, Inches(8.8), top_y, card_w, card_h, "3. Business Impact", [
        "Reduces manual dialing overhead by 90%+.",
        "Zero agent idle time with automated audio playback.",
        "Accurate call outcome logging and contact status tracking.",
        "Centralized PBX integration via FusionPBX / FreeSWITCH."
    ])

    # ------------------------------------------
    # SLIDE 3: System Architecture
    # ------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, BG_LIGHT)
    add_header(s3, "Architecture", "End-to-End System & Communication Flow")

    steps = [
        ("1. Web Management Console", ["Browser-based dashboard", "Uploads contacts & set schedules", "Monitors live call queue & status"]),
        ("2. SIP / WebRTC Engine", ["SIP.js client session", "Secure WSS connection (Port 7443)", "Handles registration & signaling"]),
        ("3. FusionPBX Telephony Server", ["FreeSWITCH call routing", "Bridge out to PSTN / Gateway", "Monitors early media & answer events"]),
        ("4. Customer Call & Audio Delivery", ["Customer device rings & answers", "Automatic WAV/IVR playback starts", "Automated disconnect after broadcast"])
    ]

    for i, (stitle, spoints) in enumerate(steps):
        sx = Inches(0.9 + i * 2.95)
        add_card(s3, sx, Inches(1.8), Inches(2.75), Inches(4.9), stitle, spoints)

    # ------------------------------------------
    # SLIDE 4: Operational Workflow
    # ------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, BG_LIGHT)
    add_header(s4, "Process Flow", "Step-by-Step Dialing Execution")

    flow_items = [
        ("Step 1: Contact Ingestion", "Upload customer records via Excel template containing Name, Phone Number, Scheduled Date, and Time."),
        ("Step 2: PBX Connection", "System registers SIP extension over WebRTC and validates audio I/O devices before starting."),
        ("Step 3: Schedule Detection", "Dialer evaluates the closest scheduled call, displaying live countdown timer for upcoming calls."),
        ("Step 4: Automated Dialing", "Upon schedule arrival, system initiates outbound call to target phone number automatically."),
        ("Step 5: Voice IVR & Hangup", "When answered, audio stream broadcasts to the customer; call terminates automatically upon completion."),
        ("Step 6: Status & Log Update", "Call result (Completed, Busy, No Answer) is saved immediately into database and UI table.")
    ]

    for idx, (ftitle, fdesc) in enumerate(flow_items):
        col = idx % 2
        row = idx // 2
        fx = Inches(0.9 + col * 5.9)
        fy = Inches(1.8 + row * 1.6)

        c = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx, fy, Inches(5.6), Inches(1.4))
        c.fill.solid()
        c.fill.fore_color.rgb = WHITE
        c.line.color.rgb = BORDER_COLOR
        c.line.width = Pt(1)

        t_box = s4.shapes.add_textbox(fx + Inches(0.2), fy + Inches(0.15), Inches(5.2), Inches(1.1))
        tf = t_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = ftitle
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p.space_after = Pt(3)

        p2 = tf.add_paragraph()
        p2.text = fdesc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MAIN

    # ------------------------------------------
    # SLIDE 5: Excel Ingestion Format & Specifications
    # ------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, BG_LIGHT)
    add_header(s5, "Data Ingestion", "Excel (.xlsx) Template Structure & Field Specifications")

    table_shape = s5.shapes.add_table(5, 5, Inches(0.9), Inches(1.8), Inches(11.533), Inches(2.6))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(2.2)
    table.columns[2].width = Inches(2.0)
    table.columns[3].width = Inches(1.8)
    table.columns[4].width = Inches(3.333)

    headers = ["Field Name", "Required Format", "Example", "Mandatory", "Description"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_DARK
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.LEFT

    rows_data = [
        ["Name", "Text string", "Abebe Kebede", "Yes", "Customer or contact identifier"],
        ["Phone Number", "Digits / E.164 (+251...)", "0911223344", "Yes", "Valid routable phone number"],
        ["Date", "YYYY-MM-DD", "2026-09-16", "Yes", "Scheduled call execution date"],
        ["Time", "HH:MM (24-Hour)", "14:30", "Yes", "Scheduled call execution time"]
    ]

    for r_idx, r_data in enumerate(rows_data):
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 0 else PptxRGBColor(241, 245, 249)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_MAIN

    add_card(s5, Inches(0.9), Inches(4.7), Inches(11.533), Inches(2.0), "Parsing & Validation Standards", [
        "Standardized format ensures automatic detection by scheduling engine.",
        "Invalid phone numbers or expired timestamps are flagged in the pre-validation table prior to launch.",
        "Supports real-time additions via the 'Quick Add' interface or bulk batch imports up to thousands of records."
    ])

    # ------------------------------------------
    # SLIDE 6: Call Statuses & Telephony States
    # ------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, BG_LIGHT)
    add_header(s6, "Call Management", "Call Lifecycle Status Definitions")

    statuses = [
        ("Pending", "Initial state upon import. Call is in queue waiting for designated scheduled timestamp."),
        ("Scheduled", "Target date/time identified by the scheduler. Active countdown timer displayed."),
        ("Dialing / Ringing", "SIP INVITE dispatched to PBX. Target subscriber device is actively ringing."),
        ("Connected / In-Call", "Call answered by customer. Pre-recorded IVR message is currently playing."),
        ("Completed", "Audio broadcast delivered successfully and call released normally."),
        ("Failed / Busy", "Subscriber rejected call, line was busy, or SIP trunk failed to establish session.")
    ]

    for idx, (st_name, st_desc) in enumerate(statuses):
        col = idx % 3
        row = idx // 3
        sx = Inches(0.9 + col * 3.95)
        sy = Inches(1.8 + row * 2.5)

        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, sy, Inches(3.6), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        tb = s6.shapes.add_textbox(sx + Inches(0.2), sy + Inches(0.2), Inches(3.2), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = st_name
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = NAVY_PRIMARY
        p.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = st_desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MAIN

    # ------------------------------------------
    # SLIDE 7: SIP & PBX Integration Details
    # ------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, BG_LIGHT)
    add_header(s7, "Configuration", "PBX & Telephony Network Settings")

    add_card(s7, Inches(0.9), Inches(1.8), Inches(5.6), Inches(4.9), "FusionPBX / FreeSWITCH Parameters", [
        "WebSocket Server: wss://your-pbx-domain:7443",
        "SIP Domain / Realm: pbx.skykin.com (Configurable)",
        "SIP Port: 5060 (Signaling), 7443 (Secure WebSockets)",
        "Audio Codecs: PCMU, PCMA, Opus (G.711u / G.711a)",
        "Extension Credentials: User authentication via standard SIP secret",
        "NAT Traversal: ICE / STUN enabled for browser WebRTC compatibility"
    ])

    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.9), "Client & Browser Requirements", [
        "Supported Browsers: Chrome 90+, Edge, Firefox",
        "Microphone & Audio Permissions: Required for WebRTC media channels",
        "Auto-Play Policy: Audio context initialized on explicit user interaction",
        "Network Quality: Low latency (<150ms) to PBX server",
        "HTTPS / SSL: Valid TLS certificate required on PBX WebSocket endpoint"
    ])

    # ------------------------------------------
    # SLIDE 8: Summary & Operations Check
    # ------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, NAVY_DARK)
    add_header(s8, "Operational Readiness", "Deployment Summary & Best Practices", dark_mode=True)

    c_box = s8.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.8))
    tf_c = c_box.text_frame
    tf_c.word_wrap = True

    items = [
        ("Pre-Flight Verification", "Ensure SIP extension registers as 'Registered' before starting batch campaign."),
        ("Contact List Hygiene", "Verify phone number formats and scheduled dates/times prior to starting dialer."),
        ("Volume Control & Rate Limiting", "Set appropriate interval between calls (default: 3-5 seconds) to prevent PBX congestion."),
        ("Post-Campaign Review", "Export detailed call disposition logs for CRM reconciliation and compliance reporting.")
    ]

    for title, desc in items:
        p = tf_c.add_paragraph() if tf_c.paragraphs[0].text else tf_c.paragraphs[0]
        p.text = f"•  {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_after = Pt(2)

        p2 = tf_c.add_paragraph()
        p2.text = f"    {desc}"
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = PptxRGBColor(203, 213, 225)
        p2.space_after = Pt(14)

    for p in output_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        prs.save(p)
    print("PowerPoint presentation generated successfully.")

# ==========================================
# 2. GENERATE PROFESSIONAL DOCX (STRICTLY NO COLOR)
# ==========================================
def build_docx(output_paths):
    doc = Document()

    # Standard 1 inch margins
    for section in doc.sections:
        section.top_margin = DocxInches(1.0)
        section.bottom_margin = DocxInches(1.0)
        section.left_margin = DocxInches(1.0)
        section.right_margin = DocxInches(1.0)

    # Style definitions - strictly black text, no colors
    def set_black_font(run, size_pt=11, bold=False, italic=False):
        run.font.name = 'Calibri'
        run.font.size = DocxPt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = DocxRGBColor(0, 0, 0)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = DocxPt(14)
        p.paragraph_format.space_after = DocxPt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        set_black_font(run, size_pt=13, bold=True)
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = DocxPt(10)
        p.paragraph_format.space_after = DocxPt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        set_black_font(run, size_pt=11.5, bold=True)
        return p

    def add_body(text, bold_prefix="", bullet=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = DocxPt(2)
        p.paragraph_format.space_after = DocxPt(4)
        p.paragraph_format.line_spacing = 1.15

        if bullet:
            p.paragraph_format.left_indent = DocxInches(0.25)
            r_b = p.add_run("• ")
            set_black_font(r_b, size_pt=10.5, bold=True)

        if bold_prefix:
            r_pref = p.add_run(bold_prefix)
            set_black_font(r_pref, size_pt=10.5, bold=True)

        r_text = p.add_run(text)
        set_black_font(r_text, size_pt=10.5)
        return p

    # Document Header / Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = DocxPt(0)
    p_title.paragraph_format.space_after = DocxPt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("SKYKIN AUTOMATIC OUTBOUND DIALER")
    set_black_font(r_title, size_pt=16, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = DocxPt(14)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("System Overview, Architecture & Operational Manual\nVersion 1.0  |  September 2026")
    set_black_font(r_sub, size_pt=10, italic=True)

    # 1. Executive Summary
    add_heading_1("1. Executive Summary")
    add_body("The SkyKin Automatic Outbound Dialer is an enterprise-grade automated telephony broadcasting solution designed to execute scheduled outbound voice campaigns. Operating over WebRTC and SIP, the system automatically detects scheduled dispatch times, dials destination telephone numbers through a PBX (FusionPBX / FreeSWITCH), delivers a pre-recorded Interactive Voice Response (IVR) message upon call answer, and logs the call outcome without requiring manual agent intervention.")

    # 2. System Architecture & Prerequisites
    add_heading_1("2. System Architecture & Components")
    add_body("The platform is structured into four primary functional layers:")
    add_body("Web Management Dashboard: A clean, browser-based administrative interface for importing contact lists, configuring SIP accounts, managing call schedules, and monitoring real-time dispatch progress.", bold_prefix="•  ")
    add_body("SIP.js Signaling Engine: Client-side WebRTC signaling stack that connects to the PBX via secure WebSockets (WSS) to manage registration, INVITE sessions, and media negotiation.", bold_prefix="•  ")
    add_body("FusionPBX / FreeSWITCH Core: Telephony server handling call switching, PSTN gateway routing, early media processing, and answer state supervision.", bold_prefix="•  ")
    add_body("Audio Playback Engine: Automatic audio streaming subsystem delivering the pre-recorded voice payload immediately upon recipient answer.", bold_prefix="•  ")

    # 3. Data Ingestion & Excel File Format
    add_heading_1("3. Contact Import Specifications")
    add_body("Customer contact records can be imported in batch using standard Excel (.xlsx or .csv) files. The file structure must conform strictly to the following columns:")

    # Table for Excel Format
    tbl = doc.add_table(rows=5, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    col_widths = [DocxInches(1.5), DocxInches(1.2), DocxInches(1.3), DocxInches(2.5)]
    headers = ["Column Header", "Format", "Required", "Description"]

    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_black_font(p.runs[0], size_pt=10, bold=True)

    row_data = [
        ("Name", "Text string", "Yes", "Customer or subscriber full name"),
        ("Phone Number", "Digits / E.164", "Yes", "Destination phone number (e.g. 0911223344)"),
        ("Date", "YYYY-MM-DD", "Yes", "Scheduled date of execution"),
        ("Time", "HH:MM (24-hr)", "Yes", "Scheduled time of execution")
    ]

    for row_idx, data in enumerate(row_data):
        cells = tbl.rows[row_idx + 1].cells
        for col_idx, text_val in enumerate(data):
            cells[col_idx].text = text_val
            p = cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_black_font(p.runs[0], size_pt=9.5)

    # Set table borders to clean black/gray
    tbl_pr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '999999')
        borders.append(border)
    tbl_pr.append(borders)

    # 4. Operational Workflow
    add_heading_1("4. Operational Workflow & Execution Steps")
    add_body("1. System Initialization: Open the dialer dashboard and verify that the SIP registration status indicates 'Registered'.")
    add_body("2. Contact List Upload: Select 'Upload Excel' to load the campaign list. The system validates the schema and loads records into the dispatch table.")
    add_body("3. Schedule Detection: The dialer analyzes all pending records, detects the closest scheduled call, and presents a countdown timer.")
    add_body("4. Automatic Dialing: When the scheduled timestamp arrives, the system executes an automated SIP call to the contact.")
    add_body("5. Call Processing: Upon answer, the pre-recorded IVR audio message plays to completion; the call terminates automatically.")
    add_body("6. Logging & Status Update: The call outcome is recorded with precise timestamps in the local database and log view.")

    # 5. Call Status Reference Guide
    add_heading_1("5. Call Status Reference")
    add_body("Pending: Contact loaded into campaign queue; awaiting scheduled trigger time.", bold_prefix="•  ")
    add_body("Scheduled: Identified as the immediate upcoming call with active countdown timer.", bold_prefix="•  ")
    add_body("Dialing / Ringing: Outbound SIP INVITE initiated; remote subscriber endpoint is ringing.", bold_prefix="•  ")
    add_body("Connected / In-Call: Call answered; automated voice playback is actively streaming.", bold_prefix="•  ")
    add_body("Completed: Call concluded successfully after full broadcast delivery.", bold_prefix="•  ")
    add_body("Failed / Busy: Call was rejected, subscriber was busy, or network error occurred.", bold_prefix="•  ")

    # 6. Technical Configuration Reference
    add_heading_1("6. Technical Configuration Parameters")
    add_body("WSS Server URL: Secure WebSocket address of the PBX server (e.g. wss://pbx.domain:7443).", bold_prefix="•  ")
    add_body("SIP Domain: PBX realm configured for extension registration.", bold_prefix="•  ")
    add_body("SIP Extension / Secret: Dedicated extension credentials provisioned on FusionPBX.", bold_prefix="•  ")
    add_body("Call Delay Interval: Configurable rest interval (3–5 seconds) between calls to allow trunk release.", bold_prefix="•  ")

    # 7. Troubleshooting Guide
    add_heading_1("7. Troubleshooting & Error Resolution")
    add_body("SIP Registration Fails: Check network connectivity, verify WSS port 7443 is open, and confirm SSL certificate validity on PBX.", bold_prefix="•  ")
    add_body("No Audio on Call Answer: Check browser microphone and audio output permissions; verify WebRTC codec compatibility (PCMU/PCMA).", bold_prefix="•  ")
    add_body("Calls Not Triggering on Time: Ensure client system clock is synchronized and date/time format matches YYYY-MM-DD HH:MM.", bold_prefix="•  ")

    for p in output_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        doc.save(p)
    print("Word documentation generated successfully.")

if __name__ == "__main__":
    desktop_dir = r"C:\Users\user\Desktop"
    project_dir = r"c:\Users\user\Desktop\call center test\automatic dialer"

    pptx_paths = [
        os.path.join(desktop_dir, "SkyKin_AutoDialer_Presentation.pptx"),
        os.path.join(project_dir, "SkyKin_AutoDialer_Presentation.pptx")
    ]
    docx_paths = [
        os.path.join(desktop_dir, "SkyKin_AutoDialer_Documentation.docx"),
        os.path.join(project_dir, "SkyKin_AutoDialer_Documentation.docx")
    ]

    build_pptx(pptx_paths)
    build_docx(docx_paths)
